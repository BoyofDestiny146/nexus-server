"""Knowledge-service extract / chunk tests. No Ollama, no Qdrant."""
from __future__ import annotations

from pathlib import Path

import pytest

from knowledge_service.chunk import chunk_units
from knowledge_service.extract import chunks_from_units, extract_chunks, extract_file
from knowledge_service.settings import settings


@pytest.fixture
def source_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    dest = tmp_path / "knowledge-sources"
    dest.mkdir()
    monkeypatch.setattr(settings, "source_dir", str(dest))
    return dest


def _write(source_dir: Path, rel: str, content: bytes) -> str:
    path = source_dir / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    return rel


def test_extract_file_returns_units_not_chunks(source_dir: Path):
    rel = _write(
        source_dir,
        "1/0/source.txt",
        b"The Bio-EV adult brief sensor measures humidity and posts a silent alert.",
    )
    extracted = extract_file("text", rel)
    assert "chunks" not in extracted
    assert extracted["unitCount"] == 1
    assert extracted["characterCount"] > 0
    assert "adult brief sensor" in extracted["units"][0]["text"]
    chunked = chunks_from_units(extracted["units"])
    assert chunked["chunkCount"] == 1
    assert chunked["chunks"][0]["contentHash"]


def test_txt_and_markdown_chunk(source_dir: Path):
    rel = _write(
        source_dir,
        "1/1/source.txt",
        b"The Bio-EV adult brief sensor measures humidity and posts a silent alert.",
    )
    out = extract_chunks("text", rel)
    assert out["chunkCount"] == 1
    assert "adult brief sensor" in out["chunks"][0]["text"]
    assert out["chunks"][0]["contentHash"]
    assert out["metadata"]["parser"] == "utf8"

    md = _write(source_dir, "1/2/source.md", b"# Overview\n\nHumidity posts a silent alert.")
    md_out = extract_chunks("markdown", md)
    assert md_out["chunkCount"] >= 1


def test_pdf_by_page(source_dir: Path):
    from pypdf import PdfWriter

    rel = "1/3/source.pdf"
    path = source_dir / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    writer = PdfWriter()
    writer.add_blank_page(width=400, height=400)
    writer.add_blank_page(width=400, height=400)
    with path.open("wb") as fh:
        writer.write(fh)
    out = extract_chunks("pdf", rel)
    assert out["metadata"]["parser"] == "pypdf"
    assert out["metadata"]["pages"] == 2


def test_chunk_units_keep_page_numbers():
    chunks = chunk_units(
        [
            {
                "text": "Page one explains the adult brief sensor.",
                "pageNumber": 1,
                "slideNumber": None,
                "sectionTitle": "Page 1",
            }
        ]
    )
    assert chunks[0]["pageNumber"] == 1


def test_docx_headings_and_tables(source_dir: Path):
    from docx import Document

    doc = Document()
    doc.add_heading("Adult Brief Sensor", level=1)
    doc.add_paragraph("The sensor measures humidity.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Alert"
    table.cell(1, 1).text = "Silent"
    rel = "1/4/source.docx"
    path = source_dir / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    out = extract_chunks("docx", rel)
    texts = " ".join(c["text"] for c in out["chunks"])
    assert "humidity" in texts.lower()
    assert "Silent" in texts
    assert any(c.get("sectionTitle") == "Adult Brief Sensor" for c in out["chunks"])


def test_pptx_slide_text_and_notes(source_dir: Path):
    from pptx import Presentation
    from pptx.util import Inches

    pres = Presentation()
    layout = pres.slide_layouts[5] if len(pres.slide_layouts) > 5 else pres.slide_layouts[0]
    slide = pres.slides.add_slide(layout)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    box.text_frame.text = "How the adult brief sensor works"
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Mention silent humidity alerts."
    rel = "1/5/source.pptx"
    path = source_dir / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    pres.save(path)
    out = extract_chunks("pptx", rel)
    assert out["chunkCount"] >= 1
    texts = " ".join(c["text"] for c in out["chunks"])
    assert all(c["slideNumber"] == 1 for c in out["chunks"])
    assert "adult brief sensor" in texts.lower()
    assert "silent humidity" in texts.lower()
    assert "Notes:" not in texts
    assert not any(c["sectionTitle"] == "Slide: 1" for c in out["chunks"])


def test_image_is_metadata_only(source_dir: Path):
    rel = _write(source_dir, "1/6/source.png", b"\x89PNG\r\n\x1a\nnot-a-real-png")
    out = extract_chunks("image", rel)
    assert out["chunkCount"] == 0
    assert out["metadata"]["skipped"] == "image"


def test_chunk_units_keep_slide_boundary():
    units = [
        {
            "text": "Slide one body " * 3,
            "pageNumber": None,
            "slideNumber": 1,
            "sectionTitle": "Slide 1",
        },
        {
            "text": "Slide two body",
            "pageNumber": None,
            "slideNumber": 2,
            "sectionTitle": "Slide 2",
        },
    ]
    chunks = chunk_units(units)
    assert {c["slideNumber"] for c in chunks} == {1, 2}
    assert [c["slideNumber"] for c in chunks] == [1, 2]


def test_pptx_title_not_flattened_and_duplicates_dropped(source_dir: Path):
    from pptx import Presentation
    from pptx.util import Inches

    pres = Presentation()
    layout = pres.slide_layouts[0]
    slide = pres.slides.add_slide(layout)
    slide.shapes.title.text = "Briefs Sensor"
    # Duplicate the title in a body-like textbox.
    box = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(8), Inches(1))
    box.text_frame.text = "Briefs Sensor"
    body = slide.shapes.add_textbox(Inches(0.5), Inches(3.0), Inches(8), Inches(2))
    body.text_frame.text = "Measures humidity and posts a silent alert."
    rel = "1/7/source.pptx"
    path = source_dir / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    pres.save(path)
    extracted = extract_file("pptx", rel)
    units = extracted["units"]
    assert units
    assert units[0]["sectionTitle"] == "Briefs Sensor"
    body_text = "\n".join(u["text"] for u in units)
    assert body_text.lower().count("briefs sensor") == 0
    assert "Slide: 2" not in body_text
    assert "Measures humidity" in body_text


def test_chunk_units_do_not_merge_unrelated_slides():
    long_a = "Adult brief sensor design. " * 40
    long_b = "Care provider benefits include quieter nights. " * 40
    chunks = chunk_units(
        [
            {"text": long_a, "slideNumber": 2, "sectionTitle": "Briefs Sensor"},
            {"text": long_b, "slideNumber": 8, "sectionTitle": "Care Provider Benefits"},
        ]
    )
    assert {c["slideNumber"] for c in chunks} == {2, 8}
    assert all(c["sectionTitle"] in {"Briefs Sensor", "Care Provider Benefits"} for c in chunks)
    # Long slides may split, but never mix the two topics in one chunk.
    for chunk in chunks:
        if chunk["slideNumber"] == 2:
            assert "care provider benefits" not in chunk["text"].lower()
        else:
            assert "adult brief sensor design" not in chunk["text"].lower()

