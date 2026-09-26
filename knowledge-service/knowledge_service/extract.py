"""Extract text units from Phase 2 source types. Images stay metadata-only."""
from __future__ import annotations

import hashlib
import logging
from pathlib import Path
from typing import Any

from fastapi import HTTPException

from .chunk import chunk_units
from .paths import absolute_source_path

log = logging.getLogger("knowledge-service")

IMAGE_TYPES = frozenset({"image"})
TEXT_TYPES = frozenset({"txt", "markdown", "text"})


def content_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    return data.decode("utf-8", errors="replace")


def extract_units(source_type: str, path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    kind = (source_type or "").strip().lower()
    if kind in IMAGE_TYPES:
        return [], {
            "parser": None,
            "skipped": "image",
            "reason": "Images are metadata-only until a vision parser is added",
            "filename": path.name,
        }
    if kind in TEXT_TYPES:
        text = _read_text(path)
        return (
            [{"text": text, "pageNumber": None, "slideNumber": None, "sectionTitle": None}],
            {"parser": "utf8", "characters": len(text)},
        )
    if kind == "pdf":
        return _extract_pdf(path)
    if kind == "docx":
        return _extract_docx(path)
    if kind == "pptx":
        return _extract_pptx(path)
    raise HTTPException(status_code=400, detail=f"unsupported sourceType: {source_type}")


def _extract_pdf(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    units: list[dict[str, Any]] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        units.append(
            {
                "text": text,
                "pageNumber": index,
                "slideNumber": None,
                "sectionTitle": f"Page {index}",
            }
        )
    return units, {"parser": "pypdf", "pages": len(reader.pages), "textPages": len(units)}


def _extract_docx(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    from docx import Document

    doc = Document(str(path))
    units: list[dict[str, Any]] = []
    heading: str | None = None
    buffer: list[str] = []
    heading_count = 0
    table_count = 0

    def flush() -> None:
        text = "\n".join(part for part in buffer if part.strip()).strip()
        buffer.clear()
        if not text:
            return
        units.append(
            {
                "text": text,
                "pageNumber": None,
                "slideNumber": None,
                "sectionTitle": heading,
            }
        )

    for para in doc.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        text = (para.text or "").strip()
        if style.startswith("Heading"):
            flush()
            heading = text or heading
            heading_count += 1
            continue
        if text:
            buffer.append(text)
    flush()

    for table in doc.tables:
        table_count += 1
        rows: list[str] = []
        for row in table.rows:
            cells = [((cell.text or "").strip()) for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                rows.append(line)
        text = "\n".join(rows).strip()
        if text:
            units.append(
                {
                    "text": text,
                    "pageNumber": None,
                    "slideNumber": None,
                    "sectionTitle": heading or "Table",
                }
            )
    return units, {
        "parser": "python-docx",
        "paragraphUnits": len(units) - table_count,
        "tables": table_count,
        "headings": heading_count,
    }


def _extract_pptx(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    pres = Presentation(str(path))
    units: list[dict[str, Any]] = []
    for index, slide in enumerate(pres.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = (shape.text or "").strip()
                if text:
                    parts.append(text)
            elif getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.TABLE:
                table = shape.table
                for row in table.rows:
                    cells = [((cell.text or "").strip()) for cell in row.cells]
                    line = " | ".join(cell for cell in cells if cell)
                    if line:
                        parts.append(line)
        notes = ""
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = (frame.text or "").strip() if frame is not None else ""
        if notes:
            parts.append(f"Notes: {notes}")
        text = "\n".join(parts).strip()
        if not text:
            continue
        units.append(
            {
                "text": text,
                "pageNumber": None,
                "slideNumber": index,
                "sectionTitle": f"Slide {index}",
            }
        )
    return units, {"parser": "python-pptx", "slides": len(list(pres.slides)), "textSlides": len(units)}


def extract_file(source_type: str, storage_path: str) -> dict[str, Any]:
    path = absolute_source_path(storage_path)
    if not path.is_file():
        raise HTTPException(status_code=404, detail="source file not found")
    units, metadata = extract_units(source_type, path)
    character_count = int(metadata.get("characters") or 0)
    if character_count <= 0:
        character_count = sum(len((unit.get("text") or "")) for unit in units)
    metadata["characterCount"] = character_count
    metadata["sourceType"] = source_type
    metadata["storagePath"] = storage_path
    metadata["unitCount"] = len(units)
    log.info(
        "extracted type=%s path=%s units=%s chars=%s",
        source_type,
        storage_path,
        len(units),
        character_count,
    )
    return {
        "units": units,
        "metadata": metadata,
        "characterCount": character_count,
        "unitCount": len(units),
    }


def chunks_from_units(units: list[dict[str, Any]]) -> dict[str, Any]:
    chunks = chunk_units(units)
    for index, chunk in enumerate(chunks):
        chunk["chunkIndex"] = index
        chunk["contentHash"] = content_hash(chunk["text"])
    return {"chunks": chunks, "chunkCount": len(chunks)}


def extract_chunks(source_type: str, storage_path: str) -> dict[str, Any]:
    extracted = extract_file(source_type, storage_path)
    chunked = chunks_from_units(list(extracted.get("units") or []))
    metadata = dict(extracted.get("metadata") or {})
    metadata["chunkCount"] = chunked["chunkCount"]
    log.info(
        "chunked type=%s path=%s chunks=%s",
        source_type,
        storage_path,
        chunked["chunkCount"],
    )
    return {
        "chunks": chunked["chunks"],
        "units": extracted.get("units") or [],
        "metadata": metadata,
        "chunkCount": chunked["chunkCount"],
        "characterCount": extracted.get("characterCount") or 0,
    }
